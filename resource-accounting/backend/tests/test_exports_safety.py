"""AUD7-SEC-04: текст из комментария/названий не должен становиться формулой в CSV/XLSX."""
import io

from openpyxl import load_workbook

from app.services.exports import BASE_COLUMNS, render_csv, render_xlsx


def _row(**over) -> dict:
    row = {key: "" for key, _ in BASE_COLUMNS}
    row.update({"meter_number": "EL-1", "consumption": "-5", "note": "=1+1"})
    row.update(over)
    return row


def test_csv_neutralizes_formula_prefix_but_keeps_numbers():
    text = render_csv([_row()]).decode("utf-8-sig")
    data_line = text.splitlines()[1]
    cells = data_line.split(";")
    keys = [key for key, _ in BASE_COLUMNS]
    assert cells[keys.index("note")] == "'=1+1"
    assert cells[keys.index("consumption")] == "-5"  # число не трогаем


def test_csv_neutralizes_other_dde_prefixes():
    for payload in ("+cmd", "-cmd", "@SUM(1)", "\tx", "\rx"):
        text = render_csv([_row(note=payload)]).decode("utf-8-sig")
        assert "'" + payload.replace("\r", "") in text.replace("\r", "") or ("'" + payload) in text


def test_xlsx_keeps_formula_text_as_string_cell():
    data = render_xlsx([_row(meter_name="=HYPERLINK(\"http://x\")")], "Ведомость")
    ws = load_workbook(io.BytesIO(data)).active
    keys = [key for key, _ in BASE_COLUMNS]
    note_cell = ws.cell(row=2, column=keys.index("note") + 1)
    name_cell = ws.cell(row=2, column=keys.index("meter_name") + 1)
    assert note_cell.data_type == "s" and note_cell.value == "=1+1"
    assert name_cell.data_type == "s" and name_cell.value.startswith("=HYPERLINK")
