"""Отчёт для проверяющих органов: GET /api/v2/requests/stats/report (.xlsx).

Данные — сид октября из `test_request_stats_analytics` (те же правила зачёта:
«выполнено» = «Принято» с `completed_at` в периоде, срок — до сдачи работы).
"""
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from uk_management_bot.database.models.apartment import Apartment
from uk_management_bot.database.models.building import Building
from uk_management_bot.database.models.request import Request
from uk_management_bot.database.models.yard import Yard

from .test_request_stats_analytics import STATS, october, people, seeded  # noqa: F401

pytestmark = pytest.mark.usefixtures("seeded")

REPORT = f"{STATS}/report"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _book(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX
    return load_workbook(io.BytesIO(resp.content))


def _rows(ws) -> list[tuple]:
    return [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]


def _indicator(ws, label: str):
    return next(r[1] for r in ws.iter_rows(values_only=True) if r[0] == label)


@pytest.mark.asyncio
async def test_report_sheets_and_summary_counts(client):
    resp = await client.get(REPORT, params=october(organization="ООО «Тест»", signer="Иванов И.И."))
    assert "official_report_2026-10-01_2026-10-31_ru.xlsx" in resp.headers["content-disposition"]
    wb = _book(resp)
    assert wb.sheetnames == ["Сводка", "Динамика", "По домам", "Реестр"]
    ws = wb["Сводка"]
    assert ws["A1"].value == "ООО «Тест»"
    assert ws["A3"].value == "за период с 01.10.2026 по 31.10.2026"
    assert _indicator(ws, "Поступило заявок") == 8
    assert _indicator(ws, "Выполнено заявок (принято жителем)") == 3
    assert _indicator(ws, "   в том числе в нормативный срок") == 2
    assert _indicator(ws, "   в том числе с нарушением срока") == 1, "high: 30 ч > 24 ч"
    assert _indicator(ws, "Доля выполненных в срок") == pytest.approx(2 / 3)
    assert _indicator(ws, "Возвращено на доработку") == 2
    # Поступившие в октябре по текущему состоянию: принято 4 (одна — уже в ноябре),
    # отменена 1, остальные в работе
    assert _indicator(ws, "   выполнено") == 4
    assert _indicator(ws, "   отменено") == 1
    assert _indicator(ws, "   в работе") == 3
    cells = [r for r in ws.iter_rows(values_only=True)]
    assert ("Сантехника", 7, 2, 1, 1) == cells[[r[0] for r in cells].index("Сантехника")][:5]
    assert any(r[3] == "/ Иванов И.И. /" for r in cells), "ФИО руководителя в подписи"
    assert len(ws._charts) == 1 and len(wb["Динамика"]._charts) == 1


@pytest.mark.asyncio
async def test_report_dynamics_by_day_in_tashkent_dates(client):
    ws = _book(await client.get(REPORT, params=october()))["Динамика"]
    rows = {r[0]: r[1:3] for r in _rows(ws) if isinstance(r[0], str) and r[0][:2].isdigit()}
    assert len(rows) == 31
    assert rows["13.10.2026"][0] == 1, "20:00 UTC 12.10 = 13.10 по Ташкенту"
    assert sum(a for _, a in rows.values()) == 3


@pytest.mark.asyncio
async def test_report_register_is_text_safe_and_ordered(client):
    ws = _book(await client.get(REPORT, params=october()))["Реестр"]
    rows = _rows(ws)
    header = next(i for i, r in enumerate(rows) if r[0] == "№")
    titles, data = rows[header], rows[header + 1:]
    assert [r[1] for r in data] == ["261001-001", "261005-001", "261003-001"], "по дате приёмки"
    first = dict(zip(titles, data[0]))
    assert isinstance(first["Дата поступления"], datetime)
    assert first["Дата поступления"] == datetime(2026, 10, 1, 10, 0), "время Ташкента"
    assert first["Выполненные работы"] == "=cmd|' /C calc'!A0"
    report_cell = ws.cell(row=next(c.row for c in ws["B"] if c.value == "261001-001"), column=7)
    assert report_cell.data_type == "s", "текст исполнителя не становится формулой"
    assert first["Срок соблюдён"] == "да"
    assert dict(zip(titles, data[2]))["Срок соблюдён"] == "нет"


@pytest.mark.asyncio
async def test_report_uzbek_labels(client):
    wb = _book(await client.get(REPORT, params=october(lang="uz")))
    assert wb.sheetnames == ["Umumiy", "Dinamika", "Uylar bo'yicha", "Reyestr"]
    assert _indicator(wb["Umumiy"], "Kelib tushgan murojaatlar") == 8


@pytest.mark.asyncio
async def test_report_by_building(client, db_session):
    yard = Yard(name="ЖК Тест")
    db_session.add(yard)
    await db_session.flush()
    building = Building(address="ул. Навои, 5", yard_id=yard.id)
    db_session.add(building)
    await db_session.flush()
    apartment = Apartment(building_id=building.id, apartment_number="12")
    db_session.add(apartment)
    await db_session.flush()
    req = await db_session.get(Request, "261001-001")
    req.apartment_id = apartment.id
    await db_session.commit()

    rows = _rows(_book(await client.get(REPORT, params=october()))["По домам"])
    by_address = {r[0]: r for r in rows}
    assert by_address["ул. Навои, 5"][1:4] == ("ЖК Тест", 1, 1)
    assert by_address["Без привязки к дому"][2:4] == (7, 2)


@pytest.mark.asyncio
async def test_report_register_masks_phones(client, db_session):
    req = await db_session.get(Request, "261001-001")
    req.description = "Течёт стояк, дом 12 кв 45, звоните +998 (90) 123-45-67"
    req.completion_report = "Заменили кран, жильца предупредили по 901234567"
    await db_session.commit()

    for lang, mark in (("ru", "[телефон]"), ("uz", "[telefon]")):
        ws = _book(await client.get(REPORT, params=october(lang=lang))).worksheets[-1]
        row = next(c.row for c in ws["B"] if c.value == "261001-001")
        assert ws.cell(row=row, column=6).value == f"Течёт стояк, дом 12 кв 45, звоните {mark}"
        assert ws.cell(row=row, column=7).value == f"Заменили кран, жильца предупредили по {mark}"


@pytest.mark.asyncio
async def test_report_empty_period(client):
    wb = _book(await client.get(REPORT, params={"date_from": "2025-01-01", "date_to": "2025-01-31"}))
    assert _indicator(wb["Сводка"], "Поступило заявок") == 0
    assert any(r[0] == "За период нет выполненных заявок." for r in _rows(wb["Реестр"]))


@pytest.mark.parametrize("params", [
    october(lang="en"),
    october(organization="x" * 201),
    {"date_from": "2026-10-31", "date_to": "2026-10-01"},
])
@pytest.mark.asyncio
async def test_report_bad_params_422(client, params):
    assert (await client.get(REPORT, params=params)).status_code == 422
