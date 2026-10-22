"""Рендер отчёта для проверяющих органов в .xlsx (openpyxl), RU/UZ.

Листы: «Сводка» (показатели, виды работ, нормативы, подписи), «Динамика»
(таблица + диаграмма), «По домам», «Реестр» (строка на выполненную заявку).
Значения — числа и даты Excel, не строки: проверяющий может сортировать и
считать сам. Текст от людей (адрес, описание, отчёт, имена) пишется строго
как текст: строка с «=» иначе станет формулой (AUD7-SEC-04, как у склада).
"""
from __future__ import annotations

import io
import math
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from uk_management_bot.utils.address_helpers import localize_address
from uk_management_bot.utils.business_time import fmt_date, fmt_datetime, to_business
from uk_management_bot.utils.categories import (
    get_category_display,
    get_urgency_display,
    resolve_category_key,
)
from uk_management_bot.utils.date_helpers import localized_month_year

from .official_report import OfficialReport, Tally

LANGUAGES = ("ru", "uz")
#: Длинные тексты в ячейке реестра обрезаются — полный текст в карточке заявки.
TEXT_MAX_CHARS = 500
#: Шапка и подписи — не больше этого (приходят из запроса).
FIELD_MAX_CHARS = 200

LABELS = {
    "ru": {
        "title": "ОТЧЁТ о выполнении заявок жителей",
        "period": "за период с {first} по {last}",
        "generated": "Сформирован: {at}",
        "sheet_summary": "Сводка",
        "sheet_dynamics": "Динамика",
        "sheet_buildings": "По домам",
        "sheet_register": "Реестр",
        "s1": "1. Основные показатели",
        "indicator": "Показатель",
        "value": "Значение",
        "created": "Поступило заявок",
        "accepted": "Выполнено заявок (принято жителем)",
        "on_time": "   в том числе в нормативный срок",
        "overdue": "   в том числе с нарушением срока",
        "on_time_share": "Доля выполненных в срок",
        "avg_hours": "Средний срок выполнения, ч",
        "returned": "Возвращено на доработку",
        "rating": "Средняя оценка жителей (1–5)",
        "ratings_count": "Число оценок",
        "cohort": "Из поступивших за период — на дату формирования:",
        "cohort_done": "   выполнено",
        "cohort_in_work": "   в работе",
        "cohort_cancelled": "   отменено",
        "s2": "2. По видам работ",
        "category": "Вид работ",
        "col_created": "Поступило",
        "col_accepted": "Выполнено",
        "col_on_time": "В срок",
        "col_overdue": "С нарушением срока",
        "col_share": "% в срок",
        "col_avg": "Средний срок, ч",
        "total": "Итого",
        "s3": "3. Соблюдение нормативных сроков",
        "urgency": "Срочность",
        "sla": "Норматив, ч",
        "notes": "Примечания",
        "note_accepted": "Выполненной считается заявка, принятая жителем в отчётном периоде.",
        "note_sla": "Срок выполнения — от поступления заявки до сдачи работы исполнителем.",
        "note_cohort": "Состояние поступивших заявок указано на дату формирования отчёта.",
        "head": "Руководитель",
        "responsible": "Ответственный исполнитель",
        "signature": "__________________",
        "stamp": "М.П.",
        "chart_dynamics": "Поступление и выполнение заявок",
        "chart_categories": "Выполнено по видам работ",
        "bucket": "Период",
        "building": "Адрес дома",
        "yard": "Двор / жилой комплекс",
        "no_building": "Без привязки к дому",
        "no": "№",
        "number": "Номер заявки",
        "created_at": "Дата поступления",
        "address": "Адрес",
        "description": "Содержание заявки",
        "report": "Выполненные работы",
        "executor": "Исполнитель",
        "executed_at": "Дата выполнения",
        "completed_at": "Дата приёмки",
        "hours": "Срок, ч",
        "sla_met": "Срок соблюдён",
        "yes": "да",
        "no_": "нет",
        "no_executor": "Не назначен",
        "truncated": "Показаны первые {shown} заявок из {total}; итоги на листе «Сводка» — по всем.",
        "empty": "За период нет выполненных заявок.",
    },
    "uz": {
        "title": "Aholi murojaatlarining bajarilishi to'g'risida HISOBOT",
        "period": "{first} dan {last} gacha bo'lgan davr uchun",
        "generated": "Tuzilgan sana: {at}",
        "sheet_summary": "Umumiy",
        "sheet_dynamics": "Dinamika",
        "sheet_buildings": "Uylar bo'yicha",
        "sheet_register": "Reyestr",
        "s1": "1. Asosiy ko'rsatkichlar",
        "indicator": "Ko'rsatkich",
        "value": "Qiymat",
        "created": "Kelib tushgan murojaatlar",
        "accepted": "Bajarilgan murojaatlar (aholi tomonidan qabul qilingan)",
        "on_time": "   shundan belgilangan muddatda",
        "overdue": "   shundan muddati buzilgan holda",
        "on_time_share": "Muddatida bajarilganlar ulushi",
        "avg_hours": "O'rtacha bajarish muddati, soat",
        "returned": "Qayta ishlashga qaytarilgan",
        "rating": "Aholining o'rtacha bahosi (1–5)",
        "ratings_count": "Baholar soni",
        "cohort": "Davr ichida kelib tushganlardan — hisobot tuzilgan sanaga:",
        "cohort_done": "   bajarilgan",
        "cohort_in_work": "   ijroda",
        "cohort_cancelled": "   bekor qilingan",
        "s2": "2. Ish turlari bo'yicha",
        "category": "Ish turi",
        "col_created": "Kelib tushgan",
        "col_accepted": "Bajarilgan",
        "col_on_time": "Muddatida",
        "col_overdue": "Muddati buzilgan",
        "col_share": "Muddatida, %",
        "col_avg": "O'rtacha muddat, soat",
        "total": "Jami",
        "s3": "3. Belgilangan muddatlarga rioya qilish",
        "urgency": "Shoshilinchlik",
        "sla": "Me'yor, soat",
        "notes": "Izohlar",
        "note_accepted": "Hisobot davrida aholi tomonidan qabul qilingan murojaat bajarilgan hisoblanadi.",
        "note_sla": "Bajarish muddati — murojaat kelib tushgandan ijrochi ishni topshirgunga qadar.",
        "note_cohort": "Kelib tushgan murojaatlarning holati hisobot tuzilgan sanaga ko'rsatilgan.",
        "head": "Rahbar",
        "responsible": "Mas'ul ijrochi",
        "signature": "__________________",
        "stamp": "M.O'.",
        "chart_dynamics": "Murojaatlarning kelib tushishi va bajarilishi",
        "chart_categories": "Ish turlari bo'yicha bajarilgan",
        "bucket": "Davr",
        "building": "Uy manzili",
        "yard": "Hovli / turar-joy majmuasi",
        "no_building": "Uyga bog'lanmagan",
        "no": "T/r",
        "number": "Murojaat raqami",
        "created_at": "Kelib tushgan sana",
        "address": "Manzil",
        "description": "Murojaat mazmuni",
        "report": "Bajarilgan ishlar",
        "executor": "Ijrochi",
        "executed_at": "Bajarilgan sana",
        "completed_at": "Qabul qilingan sana",
        "hours": "Muddat, soat",
        "sla_met": "Muddatga rioya qilingan",
        "yes": "ha",
        "no_": "yo'q",
        "no_executor": "Tayinlanmagan",
        "truncated": "Jami {total} tadan dastlabki {shown} ta murojaat ko'rsatilgan; «Umumiy» varag'idagi yakunlar — barchasi bo'yicha.",
        "empty": "Davr ichida bajarilgan murojaatlar yo'q.",
    },
}

ACCENT = "1F4E78"
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
SECTION_FONT = Font(bold=True, size=12, color=ACCENT)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATETIME_FORMAT = "DD.MM.YYYY HH:MM"
PERCENT_FORMAT = "0.0%"
HOURS_FORMAT = "0.0"
#: Высота строки по умолчанию (15 pt) в сантиметрах — для места под диаграмму.
ROW_CM = 0.53


def clean_text(value: Optional[str], limit: int = TEXT_MAX_CHARS) -> str:
    """Текст от людей: без управляющих символов (openpyxl их не пишет), с лимитом."""
    text = ILLEGAL_CHARACTERS_RE.sub("", (value or "").strip())
    if len(text) > limit:
        text = text[:limit].rstrip() + "…"
    return text


def _put(ws: Worksheet, row: int, col: int, value, *, fmt: Optional[str] = None,
         font: Optional[Font] = None, fill: Optional[PatternFill] = None,
         align: Optional[Alignment] = None, border: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if cell.data_type == "f":  # строка с «=» — текст, не формула
        cell.data_type = "s"
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = BORDER
    return cell


def _excel_dt(value: Optional[datetime]) -> Optional[datetime]:
    """Инстант → naive-время Ташкента: Excel не знает часовых поясов."""
    return to_business(value).replace(tzinfo=None) if value is not None else None


def _header_row(ws: Worksheet, row: int, titles: list[str], first_col: int = 1) -> None:
    for i, title in enumerate(titles):
        _put(ws, row, first_col + i, title, font=HEADER_FONT, fill=HEADER_FILL,
             align=CENTER, border=True)
    ws.row_dimensions[row].height = 45


def _widths(ws: Worksheet, widths: list[float]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _print_setup(ws: Worksheet, landscape: bool = False) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def _tally_cells(ws: Worksheet, row: int, col: int, t: Tally, *, created: bool = True,
                 fill: Optional[PatternFill] = None, font: Optional[Font] = None) -> int:
    values = ([t.created] if created else []) + [t.accepted, t.on_time, t.overdue]
    for value in values:
        _put(ws, row, col, value, border=True, fill=fill, font=font)
        col += 1
    _put(ws, row, col, t.on_time_share, fmt=PERCENT_FORMAT, border=True, fill=fill, font=font)
    _put(ws, row, col + 1, t.avg_hours, fmt=HOURS_FORMAT, border=True, fill=fill, font=font)
    return col + 2


def _title_block(ws: Worksheet, report: OfficialReport, L: dict, organization: str,
                 center_cols: Optional[int] = None) -> int:
    """Шапка листа; `center_cols` — центрировать по ширине таблицы, иначе
    текст слева и свободно переливается вправо (узкие листы)."""
    lines = [
        (L["title"], Font(bold=True, size=15, color=ACCENT)),
        (L["period"].format(first=fmt_date(report.first), last=fmt_date(report.last)), Font(size=12)),
        (L["generated"].format(at=fmt_datetime(report.generated_at)), Font(italic=True, color="595959")),
    ]
    if organization:
        lines.insert(0, (organization, Font(bold=True, size=12)))
    for row, (text, font) in enumerate(lines, start=1):
        _put(ws, row, 1, text, font=font,
             align=Alignment(horizontal="center") if center_cols else None)
        if center_cols:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=center_cols)
    return len(lines) + 2


def _summary(ws: Worksheet, report: OfficialReport, L: dict, lang: str,
             organization: str, signer: str) -> None:
    _widths(ws, [52, 15, 15, 15, 16, 13, 16])
    row = _title_block(ws, report, L, organization, center_cols=7)
    t = report.total

    _put(ws, row, 1, L["s1"], font=SECTION_FONT)
    row += 1
    _header_row(ws, row, [L["indicator"], L["value"]])
    row += 1
    indicators = [
        ("created", t.created, None, True),
        ("accepted", t.accepted, None, True),
        ("on_time", t.on_time, None, False),
        ("overdue", t.overdue, None, False),
        ("on_time_share", t.on_time_share, PERCENT_FORMAT, True),
        ("avg_hours", t.avg_hours, HOURS_FORMAT, False),
        ("returned", report.returned, None, False),
        ("rating", report.avg_rating, "0.0", False),
        ("ratings_count", report.ratings_count, None, False),
    ]
    for key, value, fmt, bold in indicators:
        font = Font(bold=True) if bold else None
        _put(ws, row, 1, L[key], border=True, font=font, align=WRAP)
        _put(ws, row, 2, value, fmt=fmt, border=True, font=font)
        row += 1
    _put(ws, row, 1, L["cohort"], border=True, font=Font(italic=True))
    _put(ws, row, 2, None, border=True)
    row += 1
    for key, value in (("cohort_done", report.created_done),
                       ("cohort_in_work", report.created_in_work),
                       ("cohort_cancelled", report.created_cancelled)):
        _put(ws, row, 1, L[key], border=True)
        _put(ws, row, 2, value, border=True)
        row += 1

    row += 1
    _put(ws, row, 1, L["s2"], font=SECTION_FONT)
    row += 1
    _header_row(ws, row, [L["category"], L["col_created"], L["col_accepted"], L["col_on_time"],
                          L["col_overdue"], L["col_share"], L["col_avg"]])
    row += 1
    cat_first = row
    for key, tally in report.by_category:
        _put(ws, row, 1, clean_text(get_category_display(resolve_category_key(key), lang)),
             border=True, align=WRAP)
        _tally_cells(ws, row, 2, tally)
        row += 1
    cat_last = row - 1
    _put(ws, row, 1, L["total"], border=True, font=Font(bold=True), fill=TOTAL_FILL)
    _tally_cells(ws, row, 2, t, fill=TOTAL_FILL, font=Font(bold=True))
    row += 2
    if cat_last >= cat_first:
        row = _category_chart(ws, L, cat_first, cat_last, row)

    _put(ws, row, 1, L["s3"], font=SECTION_FONT)
    row += 1
    _header_row(ws, row, [L["urgency"], L["sla"], L["col_accepted"], L["col_on_time"],
                          L["col_overdue"], L["col_share"], L["col_avg"]])
    row += 1
    for urgency, tally in report.by_urgency:
        _put(ws, row, 1, clean_text(get_urgency_display(urgency, lang)), border=True)
        _put(ws, row, 2, report.sla_hours.get(urgency), border=True)
        _tally_cells(ws, row, 3, tally, created=False)
        row += 1

    row += 1
    _put(ws, row, 1, L["notes"], font=Font(bold=True))
    for key in ("note_accepted", "note_sla", "note_cohort"):
        row += 1
        _put(ws, row, 1, f"• {L[key]}", font=Font(size=9, color="595959"))

    row += 2
    for role, name in ((L["head"], signer), (L["responsible"], "")):
        _put(ws, row, 1, role, font=Font(bold=True))
        _put(ws, row, 2, L["signature"])
        _put(ws, row, 4, f"/ {name} /" if name else "/ __________________ /")  # html-raw: ячейка xlsx, не HTML-сообщение
        row += 2
    _put(ws, row, 1, L["stamp"], font=Font(bold=True))
    _print_setup(ws)
    ws.page_setup.fitToHeight = 1  # сводка с подписями — один лист


def _category_chart(ws: Worksheet, L: dict, first: int, last: int, row: int) -> int:
    """Горизонтальные столбики «выполнено» по видам работ под таблицей;
    под диаграмму резервируются строки, чтобы следующий раздел не лёг под неё."""
    chart = BarChart()
    chart.type = "bar"
    chart.title = L["chart_categories"]
    chart.legend = None
    chart.style = 10
    chart.add_data(Reference(ws, min_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.height = max(6.0, 0.7 * (last - first + 1) + 3)
    chart.width = 22
    ws.add_chart(chart, f"A{row}")
    return row + math.ceil(chart.height / ROW_CM) + 1


def _bucket_label(key: str, granularity: str, lang: str) -> str:
    if granularity == "day":
        return fmt_date(date.fromisoformat(key))
    if granularity == "month":
        year, month = key.split("-")
        return localized_month_year(date(int(year), int(month), 1), lang)
    return key


def _dynamics(ws: Worksheet, report: OfficialReport, L: dict, lang: str, organization: str) -> None:
    _widths(ws, [22, 16, 16])
    row = _title_block(ws, report, L, organization)
    header = row
    _header_row(ws, row, [L["bucket"], L["col_created"], L["col_accepted"]])
    row += 1
    for key, created, accepted in report.series:
        _put(ws, row, 1, _bucket_label(key, report.granularity, lang), border=True)
        _put(ws, row, 2, created, border=True)
        _put(ws, row, 3, accepted, border=True)
        row += 1
    last = row - 1
    _put(ws, row, 1, L["total"], border=True, font=Font(bold=True), fill=TOTAL_FILL)
    _put(ws, row, 2, report.total.created, border=True, font=Font(bold=True), fill=TOTAL_FILL)
    _put(ws, row, 3, report.total.accepted, border=True, font=Font(bold=True), fill=TOTAL_FILL)

    chart = BarChart()
    chart.type = "col"
    chart.title = L["chart_dynamics"]
    chart.style = 10
    chart.y_axis.majorGridlines = None
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=header, max_row=last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=header + 1, max_row=last))
    chart.height = 9
    chart.width = max(18.0, min(40.0, 0.7 * len(report.series)))
    ws.add_chart(chart, "E5")
    ws.freeze_panes = ws.cell(row=header + 1, column=1)
    _print_setup(ws, landscape=True)


def _buildings(ws: Worksheet, report: OfficialReport, L: dict, lang: str, organization: str) -> None:
    _widths(ws, [44, 28, 15, 15, 15, 16, 13, 16])
    row = _title_block(ws, report, L, organization, center_cols=8)
    header = row
    _header_row(ws, row, [L["building"], L["yard"], L["col_created"], L["col_accepted"],
                          L["col_on_time"], L["col_overdue"], L["col_share"], L["col_avg"]])
    row += 1
    for b in report.by_building:
        address = clean_text(localize_address(b.address, lang)) if b.address else L["no_building"]
        _put(ws, row, 1, address, border=True, align=WRAP,
             font=None if b.address else Font(italic=True))
        _put(ws, row, 2, clean_text(b.yard), border=True, align=WRAP)
        _tally_cells(ws, row, 3, b.tally)
        row += 1
    _put(ws, row, 1, L["total"], border=True, font=Font(bold=True), fill=TOTAL_FILL)
    _put(ws, row, 2, None, border=True, fill=TOTAL_FILL)
    _tally_cells(ws, row, 3, report.total, fill=TOTAL_FILL, font=Font(bold=True))
    if row - 1 > header:
        ws.auto_filter.ref = f"A{header}:H{row - 1}"
    ws.freeze_panes = ws.cell(row=header + 1, column=1)
    _print_setup(ws, landscape=True)


def _register(ws: Worksheet, report: OfficialReport, L: dict, lang: str, organization: str) -> None:
    titles = [L["no"], L["number"], L["created_at"], L["address"], L["category"],
              L["description"], L["report"], L["urgency"], L["executor"], L["executed_at"],
              L["completed_at"], L["hours"], L["sla"], L["sla_met"]]
    _widths(ws, [6, 14, 17, 34, 18, 40, 40, 14, 22, 17, 17, 10, 11, 12])
    row = _title_block(ws, report, L, organization, center_cols=len(titles))
    if report.register_truncated:
        _put(ws, row - 1, 1, L["truncated"].format(shown=len(report.register),
                                                    total=report.total.accepted),
             font=Font(bold=True, color="C00000"))
    header = row
    _header_row(ws, row, titles)
    row += 1
    if not report.register:
        _put(ws, row, 1, L["empty"], font=Font(italic=True))
    for i, r in enumerate(report.register, start=1):
        overdue = r.is_overdue(report.sla_hours)
        values = [
            (i, None), (r.request_number, None), (_excel_dt(r.created_at), DATETIME_FORMAT),
            (clean_text(localize_address(r.address or "", lang)), None),
            (clean_text(get_category_display(resolve_category_key(r.category), lang)), None),
            (clean_text(report.descriptions.get(r.request_number)), None),
            (clean_text(r.completion_report), None),
            (clean_text(get_urgency_display(r.urgency, lang)), None),
            (clean_text(r.executor_name) or L["no_executor"], None),
            (_excel_dt(r.executed_at), DATETIME_FORMAT),
            (_excel_dt(r.completed_at), DATETIME_FORMAT),
            (r.resolution_hours, HOURS_FORMAT),
            (r.sla_limit(report.sla_hours), None),
            (None if overdue is None else (L["no_"] if overdue else L["yes"]), None),
        ]
        for col, (value, fmt) in enumerate(values, start=1):
            _put(ws, row, col, value, fmt=fmt, border=True, align=WRAP)
        if overdue:
            ws.cell(row=row, column=len(values)).font = Font(bold=True, color="C00000")
        row += 1
    if report.register:
        ws.auto_filter.ref = f"A{header}:{get_column_letter(len(titles))}{row - 1}"
    ws.freeze_panes = ws.cell(row=header + 1, column=3)
    ws.print_title_rows = f"{header}:{header}"
    _print_setup(ws, landscape=True)


def render_xlsx(report: OfficialReport, lang: str = "ru", organization: str = "",
                signer: str = "") -> bytes:
    L = LABELS[lang if lang in LANGUAGES else "ru"]
    organization = clean_text(organization, FIELD_MAX_CHARS)
    signer = clean_text(signer, FIELD_MAX_CHARS)
    wb = Workbook()
    wb.properties.title = L["title"]
    wb.properties.creator = organization or "UK Management"
    summary = wb.active
    summary.title = L["sheet_summary"]
    _summary(summary, report, L, lang, organization, signer)
    _dynamics(wb.create_sheet(L["sheet_dynamics"]), report, L, lang, organization)
    _buildings(wb.create_sheet(L["sheet_buildings"]), report, L, lang, organization)
    _register(wb.create_sheet(L["sheet_register"]), report, L, lang, organization)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_filename(first: date, last: date, lang: str) -> str:
    return f"official_report_{first.isoformat()}_{last.isoformat()}_{lang}.xlsx"
